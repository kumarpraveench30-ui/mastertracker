import io
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def export_to_csv_bytes(df: pd.DataFrame) -> bytes:
    """
    Exports DataFrame to CSV formatted bytes.
    """
    buffer = io.StringIO()
    df.to_csv(buffer, index=False, encoding='utf-8')
    return buffer.getvalue().encode('utf-8')

def export_to_excel_bytes(df: pd.DataFrame, sheet_name: str = "Filtered Data") -> bytes:
    """
    Exports DataFrame to a styled, professional Excel workbook in memory.
    """
    buffer = io.BytesIO()
    
    with pd.ExcelWriter(buffer, engine='openpyxl') as writer:
        clean_sheet = sheet_name[:31]  # Excel limits sheet names to 31 chars
        df.to_excel(writer, sheet_name=clean_sheet, index=False)
        workbook = writer.book
        worksheet = writer.sheets[clean_sheet]

        # Professional Header styling
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid") # Slate-800
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        thin_border = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1')
        )
        
        for col_idx, col_name in enumerate(df.columns, 1):
            cell = worksheet.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

            # Auto column width
            max_len = max(
                len(str(col_name)),
                df[col_name].astype(str).map(len).max() if len(df) > 0 else 0
            )
            col_letter = get_column_letter(col_idx)
            worksheet.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 40)

        # Style data cells with clean borders
        cell_font = Font(name="Segoe UI", size=10)
        for row in worksheet.iter_rows(min_row=2, max_row=len(df)+1, min_col=1, max_col=len(df.columns)):
            for cell in row:
                cell.font = cell_font
                cell.border = thin_border

    return buffer.getvalue()
